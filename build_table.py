"""Собирает data/catalog.json в таблицу Excel «Каталог_товаров.xlsx».

Листы: «Товары» (все позиции с фильтрами), «Категории» (карточки групп с сайта),
«Сводка» (количество и цены по направлениям и группам — формулами).
"""
import json
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

ROOT = Path(__file__).parent
SITE_URL = {
    "shinomontazh": "https://shinomontazh.example.com",
    "truck": "https://truck.example.com",
    "technovector": "https://technovector.example.com",
    "color-camera": "https://color-camera.example.com",
}
SITE_NAME = {
    "shinomontazh": "Легковой шиномонтаж",
    "truck": "Грузовые СТО",
    "technovector": "Сход-развал",
    "color-camera": "Малярный участок",
}
TRUCK_C = {1: "Шиномонтажные станки грузовые", 2: "Балансировочные станки грузовые", 3: "Подъёмное оборудование"}
NAME_GROUPS = [  # для позиций без явной группы — по ключевому слову в названии
    ("домкрат", "Домкраты"), ("гайковёрт", "Пневмоинструмент"), ("компрессор", "Компрессоры"),
    ("ресивер", "Ресиверы"), ("баланс", "Балансировочные станки грузовые"),
]
SKIP_SPECS = {"Группа", "Модель", "Производитель", "Диаметр обода, дюйм", "Сегмент"}

FONT = "Arial"
HEAD_FILL = PatternFill("solid", fgColor="1F3864")
ZEBRA = PatternFill("solid", fgColor="F2F5FA")
THIN = Side(style="thin", color="D0D7E2")


def group_of(item: dict, raw_by_id: dict) -> str:
    site, rid = item["id"].split(":", 1)
    if site == "truck":
        c = raw_by_id.get(("truck", rid), {}).get("c")
        if c in TRUCK_C:
            return TRUCK_C[c]
    if site == "technovector":
        return "Стенды сход-развал"
    if item["source_var"] == "CC_MODELS":
        return "Окрасочно-сушильные камеры"
    if "Группа" in item["specs"]:
        return item["specs"]["Группа"]
    low = item["name"].lower()
    return next((g for kw, g in NAME_GROUPS if kw in low), "Прочее")


def style_header(ws, ncols: int):
    for c in range(1, ncols + 1):
        cell = ws.cell(row=1, column=c)
        cell.font = Font(name=FONT, bold=True, color="FFFFFF")
        cell.fill = HEAD_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[1].height = 32
    ws.freeze_panes = "A2"


def style_body(ws, nrows: int, ncols: int, wrap_cols=()):
    for r in range(2, nrows + 2):
        for c in range(1, ncols + 1):
            cell = ws.cell(row=r, column=c)
            cell.font = Font(name=FONT, size=10)
            cell.border = Border(bottom=THIN)
            cell.alignment = Alignment(vertical="top", wrap_text=c in wrap_cols)


def add_table(ws, name: str, nrows: int, ncols: int):
    ref = f"A1:{get_column_letter(ncols)}{nrows + 1}"
    t = Table(displayName=name, ref=ref)
    t.tableStyleInfo = TableStyleInfo(name="TableStyleLight9", showRowStripes=True)
    ws.add_table(t)


def main():
    cat = json.loads((ROOT / "data/catalog.json").read_text(encoding="utf-8"))
    raw = json.loads((ROOT / "data/catalog_raw.json").read_text(encoding="utf-8"))
    raw_by_id = {}  # первая запись выигрывает: в NB_CATALOG есть группа «c», в дублях из NB_PRODUCTS — нет
    for s, vs in raw.items():
        for arr in vs.values():
            for x in arr:
                raw_by_id.setdefault((s, str(x.get("id"))), x)

    products = [x for x in cat if x["kind"] == "product"]
    categories = [x for x in cat if x["kind"] == "category"]
    for x in products:
        x["group"] = group_of(x, raw_by_id)
    products.sort(key=lambda x: (SITE_NAME[x["site"]], x["group"], x["price_rub"] or 10**12, x["name"]))

    wb = Workbook()

    # --- Товары
    ws = wb.active
    ws.title = "Товары"
    heads = ["№", "ID", "Направление", "Группа", "Наименование", "Бренд", "Модель", "Цена, ₽",
             "Статус цены", "Питание, В", "Диаметр обода, дюйм", "Сегмент", "Характеристики", "Источник"]
    ws.append(heads)
    for i, x in enumerate(products, 1):
        s = x["specs"]
        ws.append([
            i, x["id"], SITE_NAME[x["site"]], x["group"], x["name"], x["brand"] or "", s.get("Модель", ""),
            x["price_rub"], "есть" if x["price_rub"] else "по запросу", x["voltage"] or "",
            s.get("Диаметр обода, дюйм", ""), s.get("Сегмент", ""),
            "; ".join(f"{k}: {v}" for k, v in s.items() if k not in SKIP_SPECS and v),
            SITE_URL[x["site"]],
        ])
    n = len(products)
    style_header(ws, len(heads))
    style_body(ws, n, len(heads), wrap_cols={5, 13})
    for r in range(2, n + 2):
        ws.cell(row=r, column=8).number_format = '#,##0 "₽";;"—"'
        ws.cell(row=r, column=8).alignment = Alignment(horizontal="right", vertical="top")
        ws.cell(row=r, column=1).alignment = Alignment(horizontal="center", vertical="top")
    for col, w in zip("ABCDEFGHIJKLMN", [6, 22, 18, 30, 60, 14, 18, 14, 12, 10, 12, 14, 70, 30]):
        ws.column_dimensions[col].width = w
    add_table(ws, "Товары", n, len(heads))

    # --- Категории
    wc = wb.create_sheet("Категории")
    ch = ["ID", "Направление", "Карточка категории", "На сайте", "Комментарий"]
    wc.append(ch)
    for x in categories:
        r = raw_by_id.get((x["site"], x["id"].split(":", 1)[1]), {})
        wc.append([x["id"], SITE_NAME[x["site"]], x["name"],
                   " · ".join(v for v in (r.get("stock"), r.get("special"), r.get("credit")) if v),
                   "Группа товаров, а не позиция: в КП не включать"])
    style_header(wc, len(ch))
    style_body(wc, len(categories), len(ch), wrap_cols={3, 4})
    for col, w in zip("ABCDE", [24, 20, 45, 50, 40]):
        wc.column_dimensions[col].width = w

    # --- Сводка (формулы по листу «Товары»)
    wsum = wb.create_sheet("Сводка", 0)
    wsum["A1"] = "Каталог АО «Автосервис-Снаб» — сводка"
    wsum["A1"].font = Font(name=FONT, bold=True, size=14)
    wsum["A2"] = ("Источник: публичные страницы shinomontazh / truck / technovector / color-camera.example.com "
                  "(данные встроены в HTML), собрано parse_catalog.py. Это витрина сайта, а не PIM: "
                  "на сайте заявлено 21 738 позиций в прайсе.")
    wsum["A2"].font = Font(name=FONT, size=9, italic=True, color="555555")
    wsum.merge_cells("A2:F2")
    wsum["A2"].alignment = Alignment(wrap_text=True, vertical="top")
    wsum.row_dimensions[2].height = 40

    sh = ["Направление", "Группа", "Позиций", "С ценой", "Мин. цена, ₽", "Макс. цена, ₽"]
    for c, h in enumerate(sh, 1):
        cell = wsum.cell(row=4, column=c, value=h)
        cell.font = Font(name=FONT, bold=True, color="FFFFFF")
        cell.fill = HEAD_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    pairs = sorted({(SITE_NAME[x["site"]], x["group"]) for x in products})
    rng = lambda col: f"Товары!${col}$2:${col}${n + 1}"  # noqa: E731
    r = 5
    for site, grp in pairs:
        crit = f'{rng("C")},$A{r},{rng("D")},$B{r}'
        wsum.append([site, grp,
                     f"=COUNTIFS({crit})",
                     f'=COUNTIFS({crit},{rng("H")},">0")',
                     f'=IF(D{r}=0,"—",_xlfn.MINIFS({rng("H")},{crit},{rng("H")},">0"))',
                     f'=IF(D{r}=0,"—",_xlfn.MAXIFS({rng("H")},{crit}))'])
        r += 1
    last = r - 1
    wsum.append(["Итого", "", f"=SUM(C5:C{last})", f"=SUM(D5:D{last})",
                 f'=_xlfn.MINIFS({rng("H")},{rng("H")},">0")', f'=MAX({rng("H")})'])
    for row in wsum.iter_rows(min_row=5, max_row=r, max_col=6):
        for cell in row:
            cell.font = Font(name=FONT, size=10, bold=cell.row == r)
            cell.border = Border(bottom=THIN)
            if cell.row % 2 == 0 and cell.row != r:
                cell.fill = ZEBRA
            if cell.column >= 5:
                cell.number_format = '#,##0 "₽"'
                cell.alignment = Alignment(horizontal="right")
    wsum.cell(row=r + 2, column=1, value="Карточек категорий (не товары), лист «Категории»:").font = Font(name=FONT, size=10)
    wsum.cell(row=r + 2, column=3, value=f"=COUNTA(Категории!A2:A{len(categories) + 1})").font = Font(name=FONT, size=10)
    for col, w in zip("ABCDEF", [22, 36, 11, 11, 15, 15]):
        wsum.column_dimensions[col].width = w
    wsum.freeze_panes = "A5"

    out = ROOT / "Каталог_товаров.xlsx"
    wb.save(out)
    print(f"{out}: товаров {n}, категорий {len(categories)}, строк сводки {len(pairs)}")


if __name__ == "__main__":
    main()
